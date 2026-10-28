import docx
from docx.shared import Inches

doc = docx.Document()

doc.add_heading('Exploratory Data Analysis of Automotive Trends', 0)

doc.add_heading('1. Executive Summary & Objective', level=1)
doc.add_paragraph(
    "This report performs exploratory data analysis (EDA) on an automotive trends dataset containing "
    "600 simulated sales records from 2022 to 2024. The objective is to identify market demand, "
    "pricing trends, regional preferences, and clean corrupt data outliers."
)

doc.add_heading('2. Data Cleaning Methodology', level=1)
doc.add_paragraph(
    "1. Missing Values: Imputed missing prices using median pricing per vehicle category, "
    "and rating missing values using average satisfaction rating.\n"
    "2. Outliers: Filtered negative prices and extreme pricing values using 1.5 * IQR thresholds."
)

doc.add_heading('3. Key Visualizations & Findings', level=1)

images = [
    ('fig1_sales_by_vehicle_type.png', 'Figure 1: Total Units Sold by Vehicle Type'),
    ('fig2_monthly_sales_trend.png', 'Figure 2: Monthly Sales Volume Trend'),
    ('fig3_price_distribution.png', 'Figure 3: Vehicle Price Distribution'),
    ('fig4_price_vs_units_sold.png', 'Figure 4: Vehicle Price vs Units Sold'),
    ('fig5_regional_sales_heatmap.png', 'Figure 5: Regional Sales Heatmap')
]

for img_path, caption in images:
    doc.add_heading(caption, level=2)
    doc.add_picture(img_path, width=Inches(5.5))
    doc.add_paragraph("Visual evidence generated during data analysis.")

doc.add_heading('4. Business Recommendations', level=1)
doc.add_paragraph(
    "1. Expand EV & SUV inventory in top-performing regions.\n"
    "2. Optimize localized price adjustments based on median purchasing power."
)

doc.save("Automotive_Trends_EDA_Report.docx")
print("Report generated: Automotive_Trends_EDA_Report.docx")